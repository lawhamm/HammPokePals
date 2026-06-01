#!/usr/bin/env python3
"""Extract the reference catalogue (moves, abilities, items) from a PTE
Character Sheet workbook into seed-content/reference.json, which the content
loader ingests into the reference_entries table.

Usage:
    python3 scripts/xlsx-to-reference.py "/path/to/PTE_Character_Sheet.xlsx"

Requires openpyxl. Reads three sheets:
  - "Attack Data"  -> moves     (Name, Type, Class, Frequency, Range, AC, DB, Effect, ...)
  - "Ability Data" -> abilities (Name, Frequency, Base Effect, Trigger, Target, Keywords, Effect)
  - "Item Data"    -> items     (paired Name/Effect columns per category)
"""
import json
import sys
from pathlib import Path

import openpyxl

BLANKS = {None, "", "--", "None", "N/A"}


def clean(v):
    return "" if v is None else str(v).strip()


def real(v):
    s = clean(v)
    return "" if s in BLANKS else s


def shorten(text, n=160):
    text = " ".join(clean(text).split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def header_index(rows):
    """Find the header row (first row whose first cell is a non-empty label)."""
    for i, r in enumerate(rows):
        if r and real(r[0]):
            return i
    return 0


def load_moves(wb):
    ws = wb["Attack Data"]
    rows = list(ws.iter_rows(values_only=True))
    hi = header_index(rows)
    hdr = [clean(c) for c in rows[hi]]
    # Map by header label so column order changes don't break us.
    def col(label):
        return hdr.index(label) if label in hdr else None

    idx = {k: col(k) for k in [
        "Attack Name", "Type", "Class", "Frequency", "Range", "AC", "DB",
        "Effect", "Versatile Effect", "Attack Tier",
    ]}
    out = []
    for r in rows[hi + 1:]:
        name = real(r[idx["Attack Name"]]) if idx["Attack Name"] is not None else ""
        if not name:
            continue
        get = lambda key: real(r[idx[key]]) if idx[key] is not None and idx[key] < len(r) else ""
        data = {
            "Type": get("Type"),
            "Class": get("Class"),
            "Frequency": get("Frequency"),
            "Range": get("Range"),
            "Accuracy Check": get("AC"),
            "Damage Base": get("DB"),
            "Effect": get("Effect"),
            "Versatile Effect": get("Versatile Effect"),
            "Tier": get("Attack Tier"),
        }
        data = {k: v for k, v in data.items() if v}
        bits = [b for b in (get("Class"), get("Frequency")) if b]
        out.append({
            "kind": "move",
            "name": name,
            "category": get("Type"),
            "summary": " · ".join(bits) or shorten(get("Effect")),
            "data": data,
        })
    return out


def load_abilities(wb):
    ws = wb["Ability Data"]
    rows = list(ws.iter_rows(values_only=True))
    hi = header_index(rows)
    hdr = [clean(c) for c in rows[hi]]
    def col(label):
        return hdr.index(label) if label in hdr else None

    idx = {k: col(k) for k in
           ["Name", "Frequency", "Base Effect", "Trigger", "Target", "Keywords", "Effect"]}
    out = []
    for r in rows[hi + 1:]:
        name = real(r[idx["Name"]]) if idx["Name"] is not None else ""
        if not name:
            continue
        get = lambda key: real(r[idx[key]]) if idx[key] is not None and idx[key] < len(r) else ""
        data = {
            "Frequency": get("Frequency"),
            "Trigger": get("Trigger"),
            "Target": get("Target"),
            "Keywords": get("Keywords"),
            "Effect": get("Effect") or get("Base Effect"),
        }
        data = {k: v for k, v in data.items() if v}
        out.append({
            "kind": "ability",
            "name": name,
            "category": get("Keywords"),
            "summary": shorten(get("Base Effect") or get("Effect")),
            "data": data,
        })
    return out


def load_items(wb):
    ws = wb["Item Data"]
    rows = list(ws.iter_rows(values_only=True))
    hi = header_index(rows)
    # Category lives in the name-column header; each category is an independent
    # vertical list of (name, effect) pairs. Pairs: (0,1),(2,3),(4,5),(6,7),(8,9),(10,11).
    # Skip the (12,13) "All Items"/"All Effects" union to avoid duplicates.
    hdr = [clean(c) for c in rows[hi]]
    pairs = [(0, 1), (2, 3), (4, 5), (6, 7), (8, 9), (10, 11)]
    out = []
    seen = set()
    for name_col, eff_col in pairs:
        if name_col >= len(hdr):
            continue
        category = hdr[name_col].replace(" List", "").strip() or "Item"
        for r in rows[hi + 1:]:
            if name_col >= len(r):
                continue
            name = real(r[name_col])
            if not name:
                continue
            effect = real(r[eff_col]) if eff_col < len(r) else ""
            key = (name.lower(), category.lower())
            if key in seen:
                continue
            seen.add(key)
            data = {"Category": category}
            if effect:
                data["Effect"] = effect
            out.append({
                "kind": "item",
                "name": name,
                "category": category,
                "summary": shorten(effect),
                "data": data,
            })
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python3 scripts/xlsx-to-reference.py <workbook.xlsx>")
    src = Path(sys.argv[1])
    out_path = Path(__file__).resolve().parent.parent / "seed-content" / "reference.json"

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    entries = []
    counts = {}
    for name, loader in (("Attack Data", load_moves), ("Ability Data", load_abilities),
                         ("Item Data", load_items)):
        if name in wb.sheetnames:
            part = loader(wb)
            counts[name] = len(part)
            entries.extend(part)

    out_path.write_text(json.dumps(entries, ensure_ascii=False, indent=0), encoding="utf-8")
    print(f"Wrote {len(entries)} reference entries to {out_path}")
    for k, v in counts.items():
        print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
